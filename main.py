import requests
import openpyxl
from bs4 import BeautifulSoup
import re
import time
import concurrent.futures

# Load Excel file
wbkName = 'Emails.xlsx'
wbk = openpyxl.load_workbook(wbkName)

# Simulate browser headers
headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
}

# Function to extract emails from a page
def extract_emails_from_url(myRow, wks):
    url_main = wks.cell(row=myRow, column=1).value  # Main URL from first column
    url_contact = wks.cell(row=myRow, column=2).value  # Contact URL from second column

    if not url_main:
        print(f"Row {myRow}: Empty main URL, skipping.")
        return

    def fetch_emails(url):
        """ Helper function to fetch emails from a given URL """
        try:
            response = requests.get(url, headers=headers, timeout=20)
            response.raise_for_status()
            soup = BeautifulSoup(response.content, 'html.parser')
            text = soup.get_text()
            emails = re.findall(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", text)
            return ', '.join(set(emails)) if emails else None
        except requests.exceptions.RequestException as e:
            print(f"Row {myRow}: Error fetching {url}: {e}")
            return "Error"
    # Check main page first
    email_main = fetch_emails(url_main)
    wks.cell(row=myRow, column=3).value = email_main if email_main else "No Email Found"

    # If no email found on main page, check contact page
    if not email_main and url_contact:
        email_contact = fetch_emails(url_contact)
        wks.cell(row=myRow, column=4).value = email_contact if email_contact else "No Email Found"
        print(f"Row {myRow}: Extracted emails from Contact Page: {email_contact}")
    else:
        wks.cell(row=myRow, column=4).value = "Not Needed"

    print(f"Row {myRow}: Extracted emails from Main Page: {email_main}")

# Function to process rows concurrently
def process_rows_concurrently(wks, start_row, end_row):
    with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
        future_to_row = {executor.submit(extract_emails_from_url, row, wks): row for row in range(start_row, end_row)}

        for future in concurrent.futures.as_completed(future_to_row):
            try:
                future.result()
            except Exception as e:
                print(f"Error processing row {future_to_row[future]}: {e}")

# Main loop to process worksheets
for wks in wbk.worksheets:
    print(f"Processing worksheet: {wks.title}")
    total_rows = 1329  # Set the total number of rows to process
    process_rows_concurrently(wks, start_row=2, end_row=total_rows + 1)
    wbk.save(wbkName)  # Save after each sheet
# Final save
wbk.save(wbkName)
print("Scraping completed. Workbook saved.")