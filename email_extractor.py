"""
Website Email Extractor

Reads website URLs from an Excel file (column A = main URL, column B = optional
contact-page URL), visits each site in parallel threads, extracts publicly listed
email addresses and writes the results back to Excel.

Usage:
    python email_extractor.py sample/websites.xlsx
    python email_extractor.py websites.xlsx --output results.xlsx --workers 10 --timeout 15
"""
import argparse
import concurrent.futures
import re
import sys

import openpyxl
import requests
from bs4 import BeautifulSoup

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}
EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
# Matches that look like emails but are really file names, e.g. logo@2x.png
JUNK_SUFFIXES = (".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".css", ".js")


def normalize_url(url):
    url = str(url).strip()
    if not url.startswith(("http://", "https://")):
        url = "https://" + url
    return url


def fetch_emails(url, timeout):
    """Return (set_of_emails, error_message_or_None) for one page."""
    try:
        response = requests.get(url, headers=HEADERS, timeout=timeout)
        response.raise_for_status()
    except requests.exceptions.RequestException as exc:
        return set(), str(exc)

    soup = BeautifulSoup(response.content, "html.parser")
    found = set(EMAIL_RE.findall(soup.get_text(" ")))
    # Emails that only appear in mailto: links are not part of the visible text
    for link in soup.select('a[href^="mailto:"]'):
        address = link["href"][len("mailto:"):].split("?")[0].strip()
        if EMAIL_RE.fullmatch(address):
            found.add(address)

    cleaned = {e.lower().rstrip(".") for e in found if not e.lower().endswith(JUNK_SUFFIXES)}
    return cleaned, None


def process_row(row_number, main_url, contact_url, timeout):
    """Check the main page first; fall back to the contact page if nothing was found."""
    emails, error = fetch_emails(normalize_url(main_url), timeout)
    source = "main page"

    if not emails and contact_url:
        emails, contact_error = fetch_emails(normalize_url(contact_url), timeout)
        source = "contact page"
        error = contact_error if not emails else None

    if emails:
        status = f"found ({source})"
    elif error:
        status = f"error: {error[:120]}"
    else:
        status = "no email found"
    return row_number, ", ".join(sorted(emails)), status


def main():
    parser = argparse.ArgumentParser(description="Extract public contact emails from a list of websites.")
    parser.add_argument("input", help="Excel file: column A = website, column B = optional contact page")
    parser.add_argument("--output", help="Output Excel file (default: <input>_results.xlsx)")
    parser.add_argument("--workers", type=int, default=5, help="Number of parallel threads (default: 5)")
    parser.add_argument("--timeout", type=int, default=20, help="Request timeout in seconds (default: 20)")
    args = parser.parse_args()

    output = args.output or args.input.rsplit(".", 1)[0] + "_results.xlsx"
    workbook = openpyxl.load_workbook(args.input)
    failed_urls = []

    for sheet in workbook.worksheets:
        sheet.cell(row=1, column=3).value = "emails"
        sheet.cell(row=1, column=4).value = "status"
        jobs = [
            (row, sheet.cell(row=row, column=1).value, sheet.cell(row=row, column=2).value)
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value
        ]
        print(f"Sheet '{sheet.title}': {len(jobs)} websites")

        # Threads only fetch pages; all Excel writes happen here in the main thread.
        with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = [executor.submit(process_row, row, main_url, contact_url, args.timeout)
                       for row, main_url, contact_url in jobs]
            for done, future in enumerate(concurrent.futures.as_completed(futures), start=1):
                row, emails, status = future.result()
                sheet.cell(row=row, column=3).value = emails
                sheet.cell(row=row, column=4).value = status
                if status.startswith("error"):
                    failed_urls.append(str(sheet.cell(row=row, column=1).value))
                print(f"[{done}/{len(jobs)}] row {row}: {status}")

    workbook.save(output)
    if failed_urls:
        with open("failed_urls.txt", "w", encoding="utf-8") as f:
            f.write("\n".join(failed_urls) + "\n")
    print(f"Done. Results saved to {output} ({len(failed_urls)} failed URLs in failed_urls.txt)")


if __name__ == "__main__":
    sys.exit(main())
